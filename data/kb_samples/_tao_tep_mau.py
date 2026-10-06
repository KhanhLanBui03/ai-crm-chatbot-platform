"""[R&D] Sinh bộ tệp mẫu kho tri thức từ nguồn trong ``_nguon/``.

Không vào image nào. Chạy bằng venv của ai-service (đã có pymupdf, python-docx, pillow,
markdown-it-py — không cài thêm gì)::

    ai-service/.venv/bin/python data/kb_samples/_tao_tep_mau.py

Nguồn nào sinh ra tệp nào:

- ``_nguon/*.md``  → PDF (5 tệp) và DOCX (4 tệp) — cần chuyển định dạng nên phải có nguồn.
- ``_nguon/huong-dan-bao-quan-noi-chien.txt`` → tệp TXT "bẩn có chủ đích": đổi sang NFD và
  thay mỗi dấu ``{ZW}`` bằng một ký tự U+200B. Giữ nguồn ở NFC cho dễ đọc, dễ sửa.
- Các tệp TXT/MD/HTML còn lại nằm thẳng ở thư mục này — chúng tự là nguồn của chính mình.
- 4 tệp lỗi sinh bằng code. Tệp > 20 MiB KHÔNG sinh ở đây (fixture test tự tạo).

Tất định: chạy lại cho ra đúng từng byte. PDF bỏ ngày tạo và mã tệp ngẫu nhiên; DOCX và ZIP
ghim mốc thời gian của từng mục trong gói.
"""

import io
import re
from datetime import datetime
import unicodedata
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from markdown_it import MarkdownIt
from PIL import Image

THU_MUC = Path(__file__).resolve().parent
NGUON = THU_MUC / "_nguon"
MOC_THOI_GIAN = (2026, 1, 1, 0, 0, 0)  # ghim cho mọi mục trong ZIP/DOCX

PDF = [
    "bang-gia-2026",
    "chinh-sach-bao-hanh",
    "huong-dan-su-dung-may-loc-nuoc",
    "chinh-sach-tra-gop",
    "quy-dinh-van-chuyen-lap-dat",
]
DOCX = [
    "chinh-sach-doi-tra",
    "chuong-trinh-khach-hang-than-thiet",
    "bang-phi-dich-vu-sua-chua",
    "quy-trinh-xu-ly-khieu-nai",
]

# CSS cho PDF: cỡ chữ và kẻ bảng như một tài liệu văn phòng bình thường.
CSS_PDF = """
body { font-size: 10.5pt; line-height: 1.35; }
h1 { font-size: 18pt; margin-bottom: 8pt; }
h2 { font-size: 14pt; margin-top: 12pt; margin-bottom: 6pt; }
h3 { font-size: 12pt; margin-top: 8pt; margin-bottom: 4pt; }
table { border-collapse: collapse; margin: 6pt 0; }
th, td { border: 1px solid #555; padding: 3pt 5pt; }
th { background-color: #e6e6e6; }
"""

TRANG_A4 = pymupdf.paper_rect("a4")
VUNG_VIET = TRANG_A4 + (56, 56, -56, -56)


def _doc(ten: str) -> str:
    return (NGUON / ten).read_text(encoding="utf-8")


# ── PDF ──────────────────────────────────────────────────────────────────────


def tao_pdf(ten: str) -> None:
    """Markdown → HTML (markdown-it) → PDF (pymupdf.Story), tự ngắt trang.

    Chuyển sang HTML trước để tiêu đề thành tiêu đề thật (cỡ chữ riêng) và bảng thành bảng
    thật — đổ thẳng Markdown vào PDF thì dấu ``#`` và ``|`` hiện ra như chữ thường.
    """
    html = MarkdownIt("commonmark").enable("table").render(_doc(f"{ten}.md"))
    story = pymupdf.Story(html=html, user_css=CSS_PDF)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    con_noi_dung = True
    while con_noi_dung:
        thiet_bi = writer.begin_page(TRANG_A4)
        con_noi_dung, _ = story.place(VUNG_VIET)
        story.draw(thiet_bi)
        writer.end_page()
    writer.close()
    _luu_pdf_tat_dinh(pymupdf.open("pdf", buf.getvalue()), THU_MUC / f"{ten}.pdf")


def _luu_pdf_tat_dinh(doc: pymupdf.Document, dich: Path) -> None:
    """Bỏ ngày tạo/sửa và không sinh mã tệp mới — chạy lại ra đúng từng byte."""
    doc.set_metadata({})
    doc.save(dich, garbage=4, deflate=True, no_new_id=True)


# ── DOCX ─────────────────────────────────────────────────────────────────────

_IN_DAM = re.compile(r"\*\*(.+?)\*\*")


def _them_doan(doan, van_ban: str) -> None:
    """Chèn một dòng vào đoạn, giữ phần ``**in đậm**`` thành run in đậm thật."""
    vi_tri = 0
    for m in _IN_DAM.finditer(van_ban):
        if m.start() > vi_tri:
            doan.add_run(van_ban[vi_tri : m.start()])
        doan.add_run(m.group(1)).bold = True
        vi_tri = m.end()
    if vi_tri < len(van_ban):
        doan.add_run(van_ban[vi_tri:])


def _them_bang(doc: Document, dong_bang: list[str]) -> None:
    hang = [
        [o.strip() for o in d.strip().strip("|").split("|")]
        for d in dong_bang
        if not re.fullmatch(r"\|?[\s:|-]+\|?", d.strip())  # bỏ dòng kẻ |---|---|
    ]
    bang = doc.add_table(rows=len(hang), cols=len(hang[0]))
    bang.style = "Table Grid"
    for i, cac_o in enumerate(hang):
        for j, o in enumerate(cac_o[: len(hang[0])]):
            bang.cell(i, j).text = ""
            _them_doan(bang.cell(i, j).paragraphs[0], o)
            if i == 0:
                for run in bang.cell(i, j).paragraphs[0].runs:
                    run.bold = True


def tao_docx(ten: str) -> None:
    """Markdown tối giản → DOCX: heading 1–3 dùng style Heading thật, danh sách dùng style
    List thật, bảng dùng ``add_table`` — để trình đọc DOCX nhận ra cấu trúc, không chỉ chữ."""
    doc = Document()
    doc.core_properties.author = "Siêu Thị Điện Máy Ngân Hà"
    doc.core_properties.title = ""
    # python-docx mặc định ghi giờ hiện tại vào core.xml — ghim lại để tất định.
    doc.core_properties.created = doc.core_properties.modified = datetime(*MOC_THOI_GIAN)

    dong_bang: list[str] = []
    for dong in _doc(f"{ten}.md").splitlines():
        if dong.lstrip().startswith("|"):
            dong_bang.append(dong)
            continue
        if dong_bang:
            _them_bang(doc, dong_bang)
            dong_bang = []
        if not dong.strip():
            continue
        if m := re.match(r"^(#{1,3})\s+(.*)", dong):
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif m := re.match(r"^\s*[-*]\s+(.*)", dong):
            _them_doan(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\s*\d+\.\s+(.*)", dong):
            _them_doan(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            _them_doan(doc.add_paragraph(), dong)
    if dong_bang:
        _them_bang(doc, dong_bang)

    buf = io.BytesIO()
    doc.save(buf)
    _ghi_zip_tat_dinh(buf.getvalue(), THU_MUC / f"{ten}.docx")


def _muc_zip(ten: str) -> zipfile.ZipInfo:
    """Mục ZIP có mốc thời gian cố định. Phải tự đặt kiểu nén: ``writestr`` nhận ``ZipInfo``
    thì dùng kiểu nén của ``ZipInfo`` (mặc định KHÔNG nén), bỏ qua cài đặt của ``ZipFile``."""
    muc = zipfile.ZipInfo(ten, MOC_THOI_GIAN)
    muc.compress_type = zipfile.ZIP_DEFLATED
    return muc


def _ghi_zip_tat_dinh(du_lieu_zip: bytes, dich: Path) -> None:
    """Ghi lại gói ZIP với mốc thời gian cố định cho mọi mục (python-docx ghi giờ hiện tại)."""
    ra = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(du_lieu_zip)) as vao, zipfile.ZipFile(
        ra, "w", zipfile.ZIP_DEFLATED
    ) as z:
        for muc in vao.infolist():
            z.writestr(_muc_zip(muc.filename), vao.read(muc.filename))
    dich.write_bytes(ra.getvalue())


def _zip_tu_noi_dung(cac_muc: dict[str, str]) -> bytes:
    ra = io.BytesIO()
    with zipfile.ZipFile(ra, "w", zipfile.ZIP_DEFLATED) as z:
        for ten, noi_dung in cac_muc.items():
            z.writestr(_muc_zip(ten), noi_dung)
    return ra.getvalue()


# ── TXT "bẩn có chủ đích" ────────────────────────────────────────────────────


def tao_txt_ban() -> None:
    """NFD + ký tự rộng bằng không: mô phỏng văn bản dán từ macOS và từ trang web."""
    van_ban = _doc("huong-dan-bao-quan-noi-chien.txt").replace("{ZW}", "​")
    dich = THU_MUC / "huong-dan-bao-quan-noi-chien.txt"
    dich.write_text(unicodedata.normalize("NFD", van_ban), encoding="utf-8")


# ── 4 tệp lỗi ────────────────────────────────────────────────────────────────


def tao_tep_loi() -> None:
    # E2 — ảnh PNG thật mang đuôi .pdf
    anh = Image.new("RGB", (200, 200), (30, 90, 160))
    buf = io.BytesIO()
    anh.save(buf, "PNG")
    (THU_MUC / "anh-doi-duoi.pdf").write_bytes(buf.getvalue())

    # E3 — ZIP thường (không có word/document.xml) mang đuôi .docx
    (THU_MUC / "nen-doi-duoi.docx").write_bytes(
        _zip_tu_noi_dung({"ghi-chu.txt": "Tệp nén thông thường, không phải tài liệu Word."})
    )

    # E4 — bảng tính Excel: đuôi không nằm trong danh sách nhận
    (THU_MUC / "bang-tinh.xlsx").write_bytes(
        _zip_tu_noi_dung(
            {
                "[Content_Types].xml": "<Types/>",
                "xl/workbook.xml": "<workbook/>",
                "xl/worksheets/sheet1.xml": "<worksheet/>",
            }
        )
    )

    # E5 — "bản scan": chụp 2 trang đầu của chính sách bảo hành thành ảnh, không còn lớp text
    goc = pymupdf.open(THU_MUC / "chinh-sach-bao-hanh.pdf")
    scan = pymupdf.open()
    for trang in list(goc)[:2]:
        anh_trang = trang.get_pixmap(dpi=150).tobytes("jpeg", jpg_quality=75)
        moi = scan.new_page(width=trang.rect.width, height=trang.rect.height)
        moi.insert_image(moi.rect, stream=anh_trang)
    _luu_pdf_tat_dinh(scan, THU_MUC / "ban-scan-bao-hanh.pdf")


def main() -> None:
    for ten in PDF:
        tao_pdf(ten)
    for ten in DOCX:
        tao_docx(ten)
    tao_txt_ban()
    tao_tep_loi()  # sau PDF: bản scan chụp từ chinh-sach-bao-hanh.pdf
    print("Đã sinh xong vào", THU_MUC)


if __name__ == "__main__":
    main()
