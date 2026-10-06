"""Trích xuất văn bản có CẤU TRÚC từ PDF/DOCX/TXT/MD/HTML — UC019 chặng 1. [PRODUCTION]

Mọi định dạng ra cùng một dạng: danh sách ``Khoi`` theo thứ tự đọc, mỗi khối là một heading,
một đoạn văn hoặc một bảng, kèm số trang nếu định dạng có trang. Bước chia đoạn
(``chia_doan.py``) chỉ làm việc với ``Khoi``, không biết tệp gốc là gì.

Heading không phải trang trí: không có heading thì UC023 không trích dẫn được "Bảng giá 2026 >
Máy lạnh" mà chỉ trích được "trang 2". Nên mỗi định dạng phải tìm ra heading theo cách của nó:

    PDF    không có thẻ heading — suy ra từ CỠ CHỮ và CHỮ ĐẬM (tệp mẫu không có mục lục)
    DOCX   style ``Heading N`` / ``Title``
    MD     dòng ``#``…``######``
    HTML   trafilatura bóc phần nội dung chính (bỏ nav/footer/script) ra markdown → như MD
    TXT    dòng ngắn VIẾT HOA TOÀN BỘ ("1. TRƯỚC LẦN SỬ DỤNG ĐẦU TIÊN")

CHẠY Ở TIẾN TRÌNH CON (``tien_trinh.py``), nên module này:
- import nhẹ — không kéo pyvi, không kéo CSDL hay cấu hình;
- chỉ ném ``AiServiceError``: ngoại lệ đi qua ranh giới tiến trình bằng pickle, và nơi gọi
  cần ``code`` để ghi ``error_message``.

KHÔNG chuẩn hoá văn bản ở đây — ``normalize_vi`` chạy ở đường ống sau khi trích xuất, cho mọi
định dạng như nhau.
"""

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from src.ai.exceptions import AiServiceError, NoTextExtractedError, ParseFailedError

LoaiKhoi = Literal["heading", "doan", "bang"]


@dataclass(frozen=True, slots=True)
class Khoi:
    """Một đơn vị nội dung theo thứ tự đọc.

    ``cap``: cấp heading (1 = lớn nhất), ``0`` với đoạn văn và bảng.
    ``trang``: số trang bắt đầu từ 1, ``None`` với định dạng không có trang (DOCX, TXT, MD, HTML).
    Bảng: mỗi dòng bảng một dòng văn bản ``ô | ô | ô``, dòng ĐẦU là dòng tiêu đề cột.
    """

    loai: LoaiKhoi
    text: str
    cap: int = 0
    trang: int | None = None


# ── PDF ──────────────────────────────────────────────────────────────────────

# Dòng bắt đầu một mục liệt kê — trong một khối PDF, xuống dòng trước ký hiệu này là xuống
# dòng THẬT; mọi xuống dòng khác là dòng tự ngắt vì hết bề ngang trang.
_DAU_MUC = re.compile(r"^\s*(?:[•●▪◦■\-–]|\d+[.)]|[a-zđ][.)])\s")

# Bảng pymupdf dò ra có nhiều "bảng rác": vùng một dòng nằm lọt trong ô tiêu đề, đường kẻ mỏng
# không có chữ. Bảng thật phải có ít nhất 2 dòng × 2 cột và có chữ.
_BANG_TOI_THIEU_DONG = 2
_BANG_TOI_THIEU_COT = 2


_GACH_CUOI_DONG = re.compile(r"-\s*\n\s*")


def _o_bang(o: str | None) -> str:
    # Ô hẹp ngắt mã sản phẩm ngay sau gạch nối ("NH-\nTL380MD") — nối liền, không chèn dấu cách,
    # nếu không mã thành "NH- TL380MD" và làn từ khoá không còn khớp "NH-TL380MD".
    return " ".join(_GACH_CUOI_DONG.sub("-", o or "").split())


def _bang_thanh_van_ban(dong_bang: list[list[str | None]]) -> str:
    cac_dong = []
    for dong in dong_bang:
        o = [_o_bang(c) for c in dong if c is not None]  # None: ô bị gộp
        if any(o):
            cac_dong.append(" | ".join(o))
    return "\n".join(cac_dong)


def _gop_dong_noi_tiep(dong_bang: list[list[str | None]]) -> list[list[str | None]]:
    """Gộp dòng NỐI TIẾP vào dòng trên — ô PDF xuống dòng bị pymupdf dò thành một dòng riêng.

    Dấu hiệu: ô đầu (cột khoá, thường là mã sản phẩm) rỗng trong khi dòng trên có ô đầu. Đo
    trên ``bang-gia-2026.pdf``: "NH-ML24I | … | lọc" rồi "| bụi mịn | | | 10 năm" — giữ tách
    thì "bụi mịn" và "10 năm" mất liên hệ với mã NH-ML24I.
    """
    ket_qua: list[list[str | None]] = []
    for dong in dong_bang:
        o = [_o_bang(c) if c is not None else None for c in dong]
        truoc = ket_qua[-1] if ket_qua else None
        if truoc and o and not o[0] and truoc[0] and len(truoc) == len(o):
            for i, c in enumerate(o):
                if c:
                    truoc[i] = f"{truoc[i]} {c}" if truoc[i] else c
            continue
        ket_qua.append(o)

    # Cột MA: chỉ có chữ ở đúng một dòng. Gặp khi glyph font dự phòng lệch vị trí trong ô tiêu
    # đề và pymupdf dò thêm cột ("Kho n ả" | "vay" cạnh "Khoản vay"). Bỏ cả cột.
    if len(ket_qua) >= 3:
        so_cot = max(len(d) for d in ket_qua)
        cot_ma = {i for i in range(so_cot) if sum(1 for d in ket_qua if i < len(d) and d[i]) <= 1}
        ket_qua = [[c for i, c in enumerate(d) if i not in cot_ma] for d in ket_qua]
    return ket_qua


def _nam_trong(ben_trong, ben_ngoai) -> bool:
    x0, y0, x1, y1 = ben_trong
    X0, Y0, X1, Y1 = ben_ngoai
    return x0 >= X0 - 1 and y0 >= Y0 - 1 and x1 <= X1 + 1 and y1 <= Y1 + 1


def _tam_trong(bbox, vung) -> bool:
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    return vung[0] <= cx <= vung[2] and vung[1] <= cy <= vung[3]


def _bang_that(trang) -> list[tuple[tuple, str]]:
    """Các bảng thật của một trang: ``(bbox, văn bản bảng)``, đã lọc bảng rác và bảng lồng."""
    ung_vien = []
    for bang in trang.find_tables().tables:
        dong_bang = _gop_dong_noi_tiep(bang.extract())
        if len(dong_bang) < _BANG_TOI_THIEU_DONG or bang.col_count < _BANG_TOI_THIEU_COT:
            continue
        van_ban = _bang_thanh_van_ban(dong_bang)
        if van_ban.count("\n") + 1 >= _BANG_TOI_THIEU_DONG:
            ung_vien.append((tuple(bang.bbox), van_ban))
    return [
        (bbox, vb)
        for i, (bbox, vb) in enumerate(ung_vien)
        if not any(j != i and _nam_trong(bbox, khac) for j, (khac, _) in enumerate(ung_vien))
    ]


def _dong_la_dam(dong: dict) -> bool:
    """Dòng ĐẬM theo đa số ký tự, không theo từng mảnh.

    Đã đo trên tệp mẫu: công cụ xuất PDF thay riêng các chữ có dấu (``ả``, ``ờ``, ``ạ``) bằng
    font dự phòng ``NotoSerif-Regular`` — KHÔNG đậm — ngay giữa heading ``CharisSIL-Bold``.
    Đòi mọi mảnh đều đậm thì gần như không heading tiếng Việt nào qua được.
    """
    dam = tong = 0
    for s in dong["spans"]:
        n = len(s["text"].strip())
        tong += n
        # flags bit 4 (16) = đậm theo pymupdf; tên font chứa "Bold" bắt các font không đặt cờ.
        if s["flags"] & 16 or "Bold" in s["font"]:
            dam += n
    return tong > 0 and dam / tong >= 0.6


def _la_heading_pdf(dong: dict, chu: str, co_than: float) -> bool:
    """Dòng ngắn, cỡ chữ lớn hơn thân bài, và (đậm, hoặc lớn hẳn — ≥ 1,3 lần cỡ thân)."""
    co = _co_dong(dong)
    if len(chu) > 200 or co <= co_than + 0.5:
        return False
    return _dong_la_dam(dong) or co >= co_than * 1.3


def _co_dong(dong: dict) -> float:
    return max((round(s["size"], 1) for s in dong["spans"] if s["text"].strip()), default=0.0)


def _phan_tich_pdf(duong_dan: Path) -> list[Khoi]:
    import pymupdf

    try:
        tai_lieu = pymupdf.open(duong_dan)
    except Exception as e:  # pymupdf ném nhiều kiểu lỗi khác nhau cho tệp hỏng
        raise ParseFailedError(f"PDF không mở được: {type(e).__name__}") from None

    with tai_lieu:
        # Lượt 1: cỡ chữ THÂN BÀI = cỡ chiếm nhiều ký tự nhất cả tài liệu. Heading là dòng đậm
        # có cỡ LỚN HƠN cỡ thân; cấp heading xếp theo thứ hạng cỡ chữ (lớn nhất = cấp 1).
        dem_co: Counter[float] = Counter()
        cac_trang = []
        for trang in tai_lieu:
            khoi_chu = trang.get_text("dict")["blocks"]
            cac_trang.append((trang, khoi_chu))
            for k in khoi_chu:
                for dong in k.get("lines", []):
                    for s in dong["spans"]:
                        dem_co[round(s["size"], 1)] += len(s["text"].strip())
        if not dem_co or sum(dem_co.values()) == 0:
            raise NoTextExtractedError(
                "PDF không có lớp chữ — có thể là bản scan, cần OCR trước khi tải lên"
            )
        co_than = dem_co.most_common(1)[0][0]

        # Lượt 2: dựng khối theo thứ tự đọc (trên xuống dưới) trong từng trang.
        muc: list[tuple[int, float, Khoi | None, str, float]] = []
        for so_trang, (trang, khoi_chu) in enumerate(cac_trang, start=1):
            bang = _bang_that(trang)
            for bbox, van_ban in bang:
                muc.append((so_trang, bbox[1], Khoi("bang", van_ban, trang=so_trang), "", 0))

            for k in khoi_chu:
                cac_dong = [
                    d
                    for d in k.get("lines", [])
                    if "".join(s["text"] for s in d["spans"]).strip()
                    and not any(_tam_trong(d["bbox"], vung) for vung, _ in bang)
                ]
                if not cac_dong:
                    continue
                # Một khối pymupdf có thể trộn heading và thân bài: tách theo từng dòng.
                dong_doan: list[str] = []
                y_doan = 0.0
                for d in cac_dong:
                    chu = "".join(s["text"] for s in d["spans"]).strip()
                    if _la_heading_pdf(d, chu, co_than):
                        if dong_doan:
                            muc.append((so_trang, y_doan, None, "\n".join(dong_doan), 0))
                            dong_doan = []
                        muc.append((so_trang, d["bbox"][1], None, chu, _co_dong(d)))
                        continue
                    if not dong_doan:
                        y_doan = d["bbox"][1]
                    if dong_doan and not _DAU_MUC.match(chu):
                        dong_doan[-1] += " " + chu  # dòng tự ngắt → nối tiếp
                    else:
                        dong_doan.append(chu)
                if dong_doan:
                    muc.append((so_trang, y_doan, None, "\n".join(dong_doan), 0))

        cac_co_heading = sorted({co for *_, co in muc if co}, reverse=True)
        cap_theo_co = {co: i + 1 for i, co in enumerate(cac_co_heading)}

        muc.sort(key=lambda m: (m[0], m[1]))
        ket_qua: list[Khoi] = []
        for so_trang, _y, khoi, chu, co in muc:
            if khoi is not None:
                ket_qua.append(khoi)
            elif co:
                truoc = ket_qua[-1] if ket_qua else None
                # Heading dài bị ngắt thành hai dòng cùng cỡ → một heading.
                if truoc and truoc.loai == "heading" and truoc.cap == cap_theo_co[co]:
                    ket_qua[-1] = Khoi("heading", f"{truoc.text} {chu}", truoc.cap, truoc.trang)
                else:
                    ket_qua.append(Khoi("heading", chu, cap_theo_co[co], so_trang))
            else:
                ket_qua.append(Khoi("doan", chu, trang=so_trang))
        return ket_qua


# ── DOCX ─────────────────────────────────────────────────────────────────────

_STYLE_HEADING = re.compile(r"^(?:Heading|Tiêu đề)\s*(\d)$", re.IGNORECASE)


def _cap_heading_docx(ten_style: str) -> int:
    if ten_style.lower() == "title":
        return 1
    khop = _STYLE_HEADING.match(ten_style)
    return int(khop.group(1)) if khop else 0


def _phan_tich_docx(duong_dan: Path) -> list[Khoi]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        tai_lieu = docx.Document(str(duong_dan))
    except Exception as e:
        raise ParseFailedError(f"DOCX không mở được: {type(e).__name__}") from None

    ket_qua: list[Khoi] = []
    # Duyệt thân tài liệu THEO THỨ TỰ: ``doc.paragraphs`` và ``doc.tables`` là hai danh sách
    # tách rời, ghép lại thì mất vị trí của bảng giữa các đoạn văn.
    for phan_tu in tai_lieu.element.body.iterchildren():
        the = phan_tu.tag.rsplit("}", 1)[-1]
        if the == "p":
            doan = Paragraph(phan_tu, tai_lieu)
            chu = doan.text.strip()
            if not chu:
                continue
            ten_style = doan.style.name if doan.style is not None else ""
            cap = _cap_heading_docx(ten_style)
            if cap:
                ket_qua.append(Khoi("heading", chu, cap))
            elif "List" in ten_style:
                ket_qua.append(Khoi("doan", f"- {chu}"))
            else:
                ket_qua.append(Khoi("doan", chu))
        elif the == "tbl":
            bang = Table(phan_tu, tai_lieu)
            dong_bang = []
            for dong in bang.rows:
                # Ô gộp ngang xuất hiện nhiều lần cùng một đối tượng — chỉ lấy một lần.
                da_thay: set[int] = set()
                o = []
                for c in dong.cells:
                    if id(c._tc) in da_thay:
                        continue
                    da_thay.add(id(c._tc))
                    o.append(c.text)
                dong_bang.append(o)
            van_ban = _bang_thanh_van_ban(dong_bang)
            if van_ban:
                ket_qua.append(Khoi("bang", van_ban))
    return ket_qua


# ── MD / HTML ────────────────────────────────────────────────────────────────

_HEADING_MD = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_DONG_KE_BANG_MD = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")
_LIEN_KET_MD = re.compile(r"!?\[([^\]]*)\]\([^)]*\)")
_NHAN_MANH_MD = re.compile(r"(\*\*|__|`)")


def _bo_dinh_dang_md(chu: str) -> str:
    """Bỏ ký hiệu markdown chỉ để trình bày; giữ chữ của liên kết/ảnh, bỏ URL."""
    return _NHAN_MANH_MD.sub("", _LIEN_KET_MD.sub(r"\1", chu)).strip()


def _dong_bang_md(dong: str) -> str:
    o = [_bo_dinh_dang_md(c) for c in dong.strip().strip("|").split("|")]
    return " | ".join(o)


def _khoi_tu_markdown(van_ban: str) -> list[Khoi]:
    ket_qua: list[Khoi] = []
    doan: list[str] = []
    bang: list[str] = []

    def dong_doan() -> None:
        if doan:
            ket_qua.append(Khoi("doan", "\n".join(doan)))
            doan.clear()

    def dong_bang() -> None:
        if bang:
            ket_qua.append(Khoi("bang", "\n".join(bang)))
            bang.clear()

    for dong in van_ban.splitlines():
        s = dong.strip()
        if s.startswith("|"):
            dong_doan()
            if not _DONG_KE_BANG_MD.match(s):
                bang.append(_dong_bang_md(s))
            continue
        dong_bang()
        if not s:
            dong_doan()
            continue
        khop = _HEADING_MD.match(s)
        if khop:
            dong_doan()
            ket_qua.append(Khoi("heading", _bo_dinh_dang_md(khop.group(2)), len(khop.group(1))))
            continue
        doan.append(_bo_dinh_dang_md(s))
    dong_doan()
    dong_bang()
    return ket_qua


def _doc_utf8(duong_dan: Path) -> str:
    # mime.nhan_dien_tep đã xác nhận UTF-8 lúc nhận tệp (UC018); utf-8-sig bỏ BOM nếu có.
    return duong_dan.read_text(encoding="utf-8-sig")


def _phan_tich_html(duong_dan: Path) -> list[Khoi]:
    import trafilatura

    # include_formatting=True: trafilatura chỉ xuất ``#`` cho heading khi bật định dạng.
    # Ký hiệu nhấn mạnh (``**``) đi kèm được _khoi_tu_markdown bỏ đi.
    md = trafilatura.extract(
        _doc_utf8(duong_dan),
        output_format="markdown",
        include_tables=True,
        include_formatting=True,
        include_links=False,
        include_images=False,
        include_comments=False,
    )
    return _khoi_tu_markdown(md or "")


# ── TXT ──────────────────────────────────────────────────────────────────────


def _la_heading_txt(dong: str) -> bool:
    """Dòng ngắn mà mọi chữ cái đều viết hoa — "I. HÌNH THỨC THANH TOÁN"."""
    chu_cai = [c for c in dong if c.isalpha()]
    return 2 <= len(chu_cai) and len(dong) <= 120 and all(c.isupper() for c in chu_cai)


def _phan_tich_txt(duong_dan: Path) -> list[Khoi]:
    ket_qua: list[Khoi] = []
    doan: list[str] = []

    def dong_doan() -> None:
        if doan:
            ket_qua.append(Khoi("doan", "\n".join(doan)))
            doan.clear()

    for dong in _doc_utf8(duong_dan).splitlines():
        s = dong.strip()
        if not s:
            dong_doan()
        elif _la_heading_txt(s):
            dong_doan()
            # Heading đầu tiên là tên tài liệu (cấp 1), các heading sau là mục (cấp 2).
            cap = 2 if any(k.loai == "heading" for k in ket_qua) else 1
            ket_qua.append(Khoi("heading", s, cap))
        else:
            doan.append(s)
    dong_doan()
    return ket_qua


# ── Điểm vào ─────────────────────────────────────────────────────────────────

_THEO_DINH_DANG = {
    "PDF": _phan_tich_pdf,
    "DOCX": _phan_tich_docx,
    "MD": lambda p: _khoi_tu_markdown(_doc_utf8(p)),
    "HTML": _phan_tich_html,
    "TXT": _phan_tich_txt,
}


def phan_tich_tep(duong_dan: str | Path, source_type: str) -> list[Khoi]:
    """Trích ``Khoi`` từ tệp ``duong_dan`` có định dạng ``source_type`` (giá trị CHECK của V202).

    Ném ``NoTextExtractedError`` khi tệp không có chữ nào (PDF scan, HTML chỉ có khung trang),
    ``ParseFailedError`` khi thư viện không đọc được tệp. Mọi lỗi khác cũng được bọc thành
    ``ParseFailedError``: ngoại lệ lạ của thư viện có thể không pickle được qua ranh giới tiến
    trình, và nơi gọi chỉ cần biết "tệp này không phân tích được".
    """
    ham = _THEO_DINH_DANG.get(source_type)
    if ham is None:
        raise ParseFailedError(f"Không có bộ phân tích cho định dạng {source_type}")
    try:
        khoi = ham(Path(duong_dan))
    except AiServiceError:
        raise
    except Exception as e:
        raise ParseFailedError(f"Lỗi khi phân tích {source_type}: {type(e).__name__}") from None

    if not any(k.text.strip() for k in khoi if k.loai != "heading"):
        raise NoTextExtractedError(f"Không trích được nội dung nào từ tệp {source_type}")
    return khoi
